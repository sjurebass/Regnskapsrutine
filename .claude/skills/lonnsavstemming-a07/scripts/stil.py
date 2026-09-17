"""Silfer-stil for Excel-rapportene. Samme uttrykk som lønnsavstemmingene fra før."""

from __future__ import annotations

import os

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MORK = "FF1B2E28"
GRONN = "FF2D5A4A"
GRA = "FF6B7280"
SVART = "FF1A1A1A"
OK_TEKST, OK_FYLL = "FF1B6E3D", "FFF0FAF4"
AVVIK_TEKST, AVVIK_FYLL = "FFC0392B", "FFFFF0F0"
FORKLART_TEKST, FORKLART_FYLL = "FFD4760A", "FFFFF8EC"

KRONER = "#,##0.00;[RED]\\-#,##0.00"
ANTALL = "#,##0;[RED]\\-#,##0"

TYNN = Side(style="thin", color="FFD1D5DB")
UNDER = Border(bottom=TYNN)

LOGO = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "silfer_logo.png")


def topp(ws, tittel: str, undertittel: str) -> None:
    ws.sheet_view.showGridLines = False
    for rad, hoyde in ((1, 19.5), (2, 25.5), (3, 15.75), (4, 6.0), (5, 3.0)):
        ws.row_dimensions[rad].height = hoyde
    ws["D2"] = tittel
    ws["D2"].font = Font(name="Arial", size=15, bold=True, color=MORK)
    ws["D3"] = undertittel
    ws["D3"].font = Font(name="Arial", size=9, color=GRA)
    if os.path.exists(LOGO):
        try:
            from openpyxl.drawing.image import Image

            bilde = Image(LOGO)
            forhold = (bilde.width or 120) / (bilde.height or 46)
            bilde.height = 46
            bilde.width = int(46 * forhold)
            ws.add_image(bilde, "A1")
        except Exception:
            pass


def bredder(ws, kart: dict[str, float]) -> None:
    for kol, b in kart.items():
        ws.column_dimensions[kol].width = b


def tabellhode(ws, rad: int, kolonner: list[tuple[str, str]]) -> None:
    """kolonner: [(tekst, 'l'|'r'), ...] fra kolonne A."""
    ws.row_dimensions[rad].height = 19.5
    for i, (tekst, just) in enumerate(kolonner, start=1):
        c = ws.cell(row=rad, column=i, value=tekst)
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=MORK)
        c.alignment = Alignment(
            vertical="center", horizontal="right" if just == "r" else "left"
        )


def seksjon(ws, rad: int, tekst: str) -> None:
    ws.row_dimensions[rad].height = 18
    c = ws.cell(row=rad, column=1, value=tekst)
    c.font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    c.fill = PatternFill("solid", fgColor=GRONN)
    c.alignment = Alignment(vertical="center")


def celle(ws, rad, kol, verdi, *, fet=False, farge=SVART, nf=None,
          just=None, ramme=True, str_=9):
    c = ws.cell(row=rad, column=kol, value=verdi)
    c.font = Font(name="Arial", size=str_, bold=fet, color=farge)
    if nf:
        c.number_format = nf
        just = just or "right"
    c.alignment = Alignment(vertical="center", horizontal=just)
    if ramme:
        c.border = UNDER
    return c


def statusregler(ws, omraade: str) -> None:
    for tekst, skrift, fyll in (
        ("OK", OK_TEKST, OK_FYLL),
        ("AVVIK", AVVIK_TEKST, AVVIK_FYLL),
        ("FORKLART", FORKLART_TEKST, FORKLART_FYLL),
        ("PERIODISERT", FORKLART_TEKST, FORKLART_FYLL),
    ):
        ws.conditional_formatting.add(
            omraade,
            CellIsRule(
                operator="equal",
                formula=[f'"{tekst}"'],
                font=Font(name="Arial", size=9, bold=True, color=skrift),
                fill=PatternFill(start_color=fyll, end_color=fyll, fill_type="solid"),
            ),
        )
