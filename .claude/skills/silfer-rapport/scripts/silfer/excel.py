"""Excel (openpyxl): Silfer-primitivene og byggeren for spec → .xlsx.

Primitivene er en utvidelse av lønnsrapportens gamle stil.py: med standardverdiene gjør de nøyaktig det samme
(samme operasjoner i samme rekkefølge), slik at kompat-skallet `silfer.stil` gir byte-like ark.
Formler skrives med engelske funksjonsnavn og komma; arknavn siteres med quote_sheetname.
"""

from __future__ import annotations

import os
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.page import PageMargins

from . import modell, tema
from .filer import os_sti
from .tekst import til_dato

MORK, GRONN, GRA, SVART, LYS, RAMME = tema.MORK, tema.GRONN, tema.GRA, tema.SVART, tema.LYS, tema.RAMME
KRONER, ANTALL = tema.KRONER, tema.ANTALL
LOGO = str(tema.LOGO)
HEADER_HOYDER = ((1, 19.5), (2, 25.5), (3, 15.75), (4, 6.0), (5, 3.0))
FORSTE_RAD = 7


# ── primitiver ────────────────────────────────────────────────────────────────

def topp(ws, tittel: str, undertittel: str, *, ncols: int | None = None, logo_hoyde: int = 46,
         bar: bool = False, logo: str | None = None, streng: bool = False) -> None:
    """Logo i A1, tittel i D2, undertittel i D3. bar=True fyller rad 5 mørk over ncols kolonner."""
    ws.sheet_view.showGridLines = False
    for rad, hoyde in HEADER_HOYDER:
        ws.row_dimensions[rad].height = hoyde
    ws["D2"] = tittel
    ws["D2"].font = Font(name="Arial", size=15, bold=True, color=MORK)
    ws["D3"] = undertittel
    ws["D3"].font = Font(name="Arial", size=9, color=GRA)
    if bar:
        for c in range(1, (ncols or 1) + 1):
            ws.cell(row=5, column=c).fill = PatternFill("solid", fgColor=MORK)
    sti = logo or LOGO
    if os.path.exists(sti):
        try:
            from openpyxl.drawing.image import Image

            bilde = Image(sti)
            forhold = (bilde.width or 120) / (bilde.height or 46)
            bilde.height = logo_hoyde
            bilde.width = int(logo_hoyde * forhold)
            ws.add_image(bilde, "A1")
        except Exception as e:
            if streng:
                raise RuntimeError(f"Logoen kunne ikke legges inn ({e}). Kjør med --with pillow.") from e
    elif streng:
        raise FileNotFoundError(sti)


def bredder(ws, kart: dict[str, float]) -> None:
    for kol, b in kart.items():
        ws.column_dimensions[kol].width = b


def tabellhode(ws, rad: int, kolonner: list[tuple[str, str]], *, hoyde: float = 19.5) -> None:
    """kolonner: [(tekst, 'l'|'r'|'c'), ...] fra kolonne A."""
    ws.row_dimensions[rad].height = hoyde
    for i, (tekst, just) in enumerate(kolonner, start=1):
        c = ws.cell(row=rad, column=i, value=tekst)
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=MORK)
        c.alignment = Alignment(
            vertical="center", horizontal="right" if just == "r" else ("center" if just == "c" else "left")
        )


def seksjon(ws, rad: int, tekst: str, *, ncols: int = 1) -> None:
    """Grønn seksjonslinje. ncols>1 fyller hele bredden."""
    ws.row_dimensions[rad].height = 18
    c = ws.cell(row=rad, column=1, value=tekst)
    c.font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    c.fill = PatternFill("solid", fgColor=GRONN)
    c.alignment = Alignment(vertical="center")
    for kol in range(2, ncols + 1):
        ws.cell(row=rad, column=kol).fill = PatternFill("solid", fgColor=GRONN)


def celle(ws, rad, kol, verdi, *, fet=False, farge=SVART, nf=None, just=None, ramme=True, str_=9,
          ramme_farge=RAMME, fyll=None, wrap=False, kursiv=False, innrykk=0):
    c = ws.cell(row=rad, column=kol, value=verdi)
    if kursiv:
        c.font = Font(name="Arial", size=str_, bold=fet, italic=True, color=farge)
    else:
        c.font = Font(name="Arial", size=str_, bold=fet, color=farge)
    if nf:
        c.number_format = nf
        just = just or "right"
    if wrap:
        c.alignment = Alignment(vertical="center", horizontal=just, wrap_text=True)
    elif innrykk:
        c.alignment = Alignment(vertical="center", horizontal=just or "left", indent=innrykk)
    else:
        c.alignment = Alignment(vertical="center", horizontal=just)
    if ramme:
        c.border = Border(bottom=Side(style="thin", color=ramme_farge))
    if fyll:
        c.fill = PatternFill("solid", fgColor=fyll)
    return c


def statusregler(ws, omraade: str, *, statuser=tema.STATUSER_ORDNET) -> None:
    """Betinget formatering per statusord, i den gitte rekkefølgen (= prioritet)."""
    for tekst in statuser:
        skrift, fyll = tema.STATUSER[tekst]
        kw = {"font": Font(name="Arial", size=9, bold=True, color=skrift)}
        if fyll:
            kw["fill"] = PatternFill(start_color=fyll, end_color=fyll, fill_type="solid")
        ws.conditional_formatting.add(omraade, CellIsRule(operator="equal", formula=[f'"{tekst}"'], **kw))


def utskrift(ws, *, landskap: bool = True, tittelrader: str | None = None) -> None:
    """A4, tilpass til bredden, bunntekst «Silfer AS · <ark>» og «Side n av N»."""
    ws.page_setup.orientation = "landscape" if landskap else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.55, bottom=0.6, header=0.3, footer=0.3)
    ws.oddFooter.left.text = f"{tema.FORFATTER} · &A"
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.font = "Arial"
    ws.oddFooter.right.text = "Side &P av &N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.font = "Arial"
    if tittelrader:
        ws.print_title_rows = tittelrader


# ── bygger ────────────────────────────────────────────────────────────────────

NF = {"belop": KRONER, "antall": ANTALL, "prosent": tema.PROSENT, "dato": tema.DATO}
JUST = {"l": "left", "r": "right", "c": "center"}


class _Plass:
    """Hvor en seksjons tabell havnet i arket."""

    def __init__(self, ark: str, kolonner: list[dict]):
        self.ark = ark
        self.kol = {k["id"]: get_column_letter(i) for i, k in enumerate(kolonner, start=1)}
        self.rad: dict[int, int] = {}       # indeks i Tabell.rader → arkrad
        self.d0 = self.d1 = 0               # første/siste datarad
        self.ib = self.ub = None
        self.hode = 0


def _excel_kolonner(t: modell.Tabell) -> list[dict]:
    return [k for k in t.kolonner if k["vis"] in ("begge", "excel")]


def _undertittel(dok: dict) -> str:
    return "   |   ".join(modell.under_deler(dok))


def _ren(tekst: str) -> str:
    """Fjerner **fet**-markering (Excel-cellen er ren tekst)."""
    return re.sub(r"\*\*(.+?)\*\*", r"\1", tekst or "")


def _tall(x) -> str:
    return f"{x:g}" if x >= 0 else f"({x:g})"


def _formel(oppskrift: tuple, r: int, p: _Plass, t: modell.Tabell, plasser: dict[str, _Plass]) -> str | None:
    L = p.kol
    art = oppskrift[0]
    try:
        if art == "lopende":
            lop = t.lopende
            s = lop["id"]
            forrige = max(rr for rr in p.rad.values() if rr < r and (rr == p.ib or p.d0 <= rr <= p.d1))
            f = f"={L[s]}{forrige}"
            if lop["lopende"]["pluss"]:
                f += f"+{L[lop['lopende']['pluss']]}{r}"
            if lop["lopende"]["minus"]:
                f += f"-{L[lop['lopende']['minus']]}{r}"
            return f
        if art == "sum":
            c = L[oppskrift[1]]
            return f"=SUM({c}{p.d0}:{c}{p.d1})"
        if art == "tell":
            c = L[oppskrift[1]]
            return f"=COUNTA({c}{p.d0}:{c}{p.d1})"
        if art == "ub_lopende":
            return f"={L[t.lopende['id']]}{p.d1 if p.d1 >= p.d0 and t.data()[-1].art == 'data' else p.ib}"
        if art == "lik":
            return f'=IF(ABS({L[oppskrift[1]]}{r}-{_tall(oppskrift[2])})<0.005,"OK","AVVIK")'
        if art == "kontroll_diff":
            lop = t.lopende
            s = L[lop["id"]]
            f = f"={s}{p.ib}"
            if lop["lopende"]["pluss"]:
                c = L[lop["lopende"]["pluss"]]
                f += f"+SUM({c}{p.d0}:{c}{p.d1})"
            if lop["lopende"]["minus"]:
                c = L[lop["lopende"]["minus"]]
                f += f"-SUM({c}{p.d0}:{c}{p.d1})"
            ub = p.ub if p.ub else p.d1
            return f + f"-{s}{ub}"
        if art == "kontroll_status":
            return f'=IF(ABS({L[oppskrift[1]]}{r})<0.005,"OK","AVVIK")'
        if art in ("countif", "sumif"):
            q = plasser[oppskrift[1]]
            ref = quote_sheetname(q.ark)
            g = q.kol[oppskrift[2]]
            omr = f"{ref}!${g}${q.d0}:${g}${q.d1}"
            if art == "countif":
                return f"=COUNTIF({omr},{L[oppskrift[2]]}{r})"
            s = q.kol[oppskrift[3]]
            return f"=SUMIF({omr},{L[oppskrift[2]]}{r},{ref}!${s}${q.d0}:${s}${q.d1})"
        if art == "avvik":
            return f"={L[oppskrift[1]]}{r}-{L[oppskrift[2]]}{r}"
        if art == "kstatus":
            return f'=IF(ABS({L[oppskrift[1]]}{r})<={oppskrift[2]:g},"OK","{oppskrift[3]}")'
        if art == "antall_avvik":
            c = L[oppskrift[1]]
            return f'=COUNTIF({c}{p.d0}:{c}{p.d1},"AVVIK")'
        if art == "pst":
            a, b = L[oppskrift[1]], L[oppskrift[2]]
            return f'=IF({b}{r}=0,"",({a}{r}-{b}{r})/ABS({b}{r})*100)'
        if art == "gsum":
            c = L[oppskrift[1]]
            return f"=SUM({c}{p.rad[oppskrift[2]]}:{c}{p.rad[oppskrift[3]]})"
        if art == "ledd":
            c = L[oppskrift[1]]
            f = ""
            for fortegn, ref in oppskrift[2]:
                ledd = f"{c}{p.rad[ref[1]]}" if ref[0] == "rad" else f"SUM({c}{p.rad[ref[1]]}:{c}{p.rad[ref[2]]})"
                f += ("-" if fortegn < 0 else ("+" if f else "")) + ledd
            return "=" + f
        if art == "lik_alle":
            vilkaar = [f"ABS({L[k]}{r}-{_tall(oppskrift[2])})<0.005" for k in oppskrift[1]]
            test = vilkaar[0] if len(vilkaar) == 1 else f"AND({','.join(vilkaar)})"
            return f'=IF({test},"OK","AVVIK")'
    except KeyError:
        return None  # en kolonne formelen trenger vises ikke i Excel → skriv verdien
    raise ValueError(f"ukjent formel {oppskrift}")


def _plasser(seksjoner: list[modell.Seksjon]) -> tuple[dict[str, list], dict[str, _Plass], dict[str, dict]]:
    """Første runde: hvilke rader hver del havner på, per ark."""
    ark: dict[str, list] = {}
    for s in seksjoner:
        ark.setdefault(s.ark, []).append(s)
    plasser, layout = {}, {}
    for navn, liste in ark.items():
        flere = len(liste) > 1
        rad = FORSTE_RAD
        for i, s in enumerate(liste):
            lay = {"bar": None, "tekst": None, "bokser": [], "punkter": [], "merknad": None}
            if i > 0:
                rad += 1
            if s.overskrift and (s.bar if s.bar is not None else flere):
                lay["bar"] = rad
                rad += 1
            if s.tekst:
                lay["tekst"] = rad
                rad += 1
            if s.tabell:
                p = _Plass(navn, _excel_kolonner(s.tabell))
                p.hode = rad
                rad += 1
                for j, r in enumerate(s.tabell.rader):
                    if r.vis == "pdf":
                        continue
                    if r.luft_for:
                        rad += 1
                    p.rad[j] = rad
                    if r.art in ("data", "tom"):
                        p.d0 = p.d0 or rad
                        p.d1 = rad
                    elif r.art == "ib":
                        p.ib = rad
                    elif r.art == "ub":
                        p.ub = rad
                    rad += 1
                plasser[s.id] = p
            for _ in s.bokser:
                rad += 1
                lay["bokser"].append(rad)
                rad += 1
            if s.punkter:
                rad += 1 if (s.tabell or s.bokser) else 0
                for _ in s.punkter:
                    lay["punkter"].append(rad)
                    rad += 1
            if s.merknad:
                rad += 1 if (s.tabell or s.bokser or s.punkter or s.tekst) else 0
                lay["merknad"] = rad
                rad += 1
            layout[s.id] = lay
    return ark, plasser, layout


def _skriv_tabell(ws, s: modell.Seksjon, p: _Plass, plasser: dict[str, _Plass]) -> list[str]:
    """Skriver tabellhode og rader. Gir statusområdene som skal ha betinget formatering."""
    t = s.tabell
    kol = _excel_kolonner(t)
    tabellhode(ws, p.hode, [(k["tekst"], k["just"]) for k in kol])
    status_omr: list[str] = []
    status_kol = [k for k in kol if k["type"] == "status"]
    for j, r in enumerate(t.rader):
        if j not in p.rad:
            continue
        rr = p.rad[j]
        zebra = t.zebra and r.art == "data" and rr % 2 == 0
        avvikrad = t.uthev_avvik and r.art == "data" and modell.avvik_i(t, r)
        for i, k in enumerate(kol, start=1):
            c_ = r.celler.get(k["id"])
            typ = (c_.type if c_ and c_.type else k["type"])
            verdi = c_.verdi if c_ else None
            formel = c_.formel if c_ else None
            if formel:
                f = _formel(formel, rr, p, t, plasser)
                if f:
                    verdi = f
            if typ == "dato":
                verdi = til_dato(verdi)
            fet = (r.fet and typ != "dato") if (c_ is None or c_.fet is None) else c_.fet
            farge = GRA if (c_ and c_.dempet) else (
                MORK if r.art in ("sum", "ub", "ib", "kontroll", "antall_avvik", "gruppe", "resultat") else SVART)
            if r.art == "tom":
                verdi, farge = None, GRA
            nf = NF.get(typ)
            if typ == "heltall" and isinstance(verdi, (int, float)):
                nf = None
            just = JUST[k["just"]] if typ == k["type"] else ("center" if typ == "status" else None)
            if typ == "antall" and k["type"] != "antall":
                just = "right"
            fyll = tema.AVVIK_FYLL if avvikrad and typ != "status" else (LYS if zebra else None)
            c = celle(ws, rr, i, verdi, fet=fet or typ == "status", farge=farge, nf=nf, just=just, fyll=fyll,
                      innrykk=c_.innrykk if c_ else 0)
            if typ == "tekst" and isinstance(verdi, str) and not (formel and verdi.startswith("=")):
                c.data_type = "s"           # tekst som starter med = + - @ blir aldri formel
            if typ == "status" and k["type"] != "status":
                status_omr.append(f"{get_column_letter(i)}{rr}")
    for k in status_kol:
        L = p.kol[k["id"]]
        if t.kilde == "oppstilling":        # statuser også på sum- og resultatlinjer
            status_omr.append(f"{L}{min(p.rad.values())}:{L}{max(p.rad.values())}")
        else:
            status_omr.append(f"{L}{p.d0}:{L}{p.d1}")
    return status_omr


def bygg_excel(spec: dict, ut) -> Path:
    dok = spec["dokument"]
    seksjoner = modell.los(spec)
    ark, plasser, layout = _plasser(seksjoner)
    wb = Workbook()
    wb.remove(wb.active)
    under = _undertittel(dok)
    for navn, liste in ark.items():
        ws = wb.create_sheet(navn)
        ncols = max([len(_excel_kolonner(s.tabell)) for s in liste if s.tabell] or [6])
        tittel = liste[0].overskrift if (len(liste) == 1 and liste[0].overskrift) else dok["tittel"]
        topp(ws, tittel, under, ncols=ncols, bar=True, streng=True)
        bredde: dict[str, float] = {}
        status_omr: list[str] = []
        frys = filter_ = tittelrader = None
        totalbredde = 0.0
        for s in liste:
            lay = layout[s.id]
            if lay["bar"]:
                seksjon(ws, lay["bar"], s.overskrift, ncols=ncols)
            if lay["tekst"]:
                celle(ws, lay["tekst"], 1, _ren(s.tekst), ramme=False)
            if s.tabell:
                p = plasser[s.id]
                kol = _excel_kolonner(s.tabell)
                for i, k in enumerate(kol, start=1):
                    L = get_column_letter(i)
                    bredde[L] = max(bredde.get(L, 0), k["bredde"])
                status_omr += _skriv_tabell(ws, s, p, plasser)
                if s.tabell.frys and frys is None:
                    frys = f"A{p.hode + 1}"
                    tittelrader = f"{p.hode}:{p.hode}"
                if s.tabell.filter and filter_ is None:
                    filter_ = f"A{p.hode}:{get_column_letter(len(kol))}{p.d1}"
            for b, rr in zip(s.bokser, lay["bokser"]):
                tekst_farge, fyll = (tema.OK_TEKST, tema.OK_FYLL) if b.type == "konkl" else (tema.FORKLART_TEKST, tema.FORKLART_FYLL)
                for kk in range(1, ncols + 1):
                    ws.cell(row=rr, column=kk).fill = PatternFill("solid", fgColor=fyll)
                celle(ws, rr, 1, _ren(b.tekst), farge=tekst_farge, fet=True, ramme=False, fyll=fyll)
                ws.row_dimensions[rr].height = 20
            for nr, (pkt, rr) in enumerate(zip(s.punkter, lay["punkter"]), start=1):
                celle(ws, rr, 1, f"{nr}. {_ren(pkt)}", ramme=False)
            if lay["merknad"]:
                m = ws.cell(row=lay["merknad"], column=1, value=_ren(s.merknad))
                m.font = Font(name="Arial", size=8, italic=True, color=GRA)
        for omr in status_omr:
            statusregler(ws, omr)
        bredder(ws, bredde)
        totalbredde = sum(bredde.values())
        if frys:
            ws.freeze_panes = frys
        if filter_:
            ws.auto_filter.ref = filter_
        utskrift(ws, landskap=totalbredde > 110, tittelrader=tittelrader)
    wb.save(os_sti(ut))
    return Path(ut)
