"""Utviklerverktøy: sammenligner to .xlsx helt ned på stil – eller bare verdiene.

    uv run --with "openpyxl==3.1.5" --with pillow python .claude/skills/silfer-rapport/scripts/sammenlign_stil.py <a.xlsx> <b.xlsx>
    … sammenlign_stil.py <a.xlsx> <b.xlsx> --verdier [--fra-rad 6]

Standard: alle deler i zip-filen unntatt docProps/* (tidsstempler) skal være byte-like; i tillegg sammenlignes celle for
celle (verdi, skrift, fyll, ramme, justering, tallformat), rad- og kolonnemål, flettede celler, betinget formatering,
bilder, frys, filter og utskrift, så en forskjell kan pekes ut. --verdier sammenligner bare verdier og formeltekst
(fra rad --fra-rad), f.eks. når stilen er endret med vilje. Exit 0 = identisk.
"""

from __future__ import annotations

import argparse
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from silfer.filer import os_sti  # noqa: E402


def zip_deler(a: str, b: str) -> list[str]:
    ut = []
    with zipfile.ZipFile(os_sti(a)) as za, zipfile.ZipFile(os_sti(b)) as zb:
        na = {n for n in za.namelist() if not n.startswith("docProps/")}
        nb = {n for n in zb.namelist() if not n.startswith("docProps/")}
        for n in sorted(na - nb):
            ut.append(f"del {n}: bare i A")
        for n in sorted(nb - na):
            ut.append(f"del {n}: bare i B")
        for n in sorted(na & nb):
            da, db = za.read(n), zb.read(n)
            if da != db:
                i = next((k for k in range(min(len(da), len(db))) if da[k] != db[k]), min(len(da), len(db)))
                ut.append(f"del {n}: ulik fra byte {i}: A …{da[max(0, i - 40):i + 60]!r} B …{db[max(0, i - 40):i + 60]!r}")
    return ut


def _farge(x):
    farge = getattr(x, "color", None) if x is not None else None
    return getattr(farge, "rgb", None)


def _stil(c) -> tuple:
    f, fy, r, j = c.font, c.fill, c.border, c.alignment
    return (
        (getattr(f, "name", None), getattr(f, "sz", None), getattr(f, "b", None), getattr(f, "i", None), _farge(f)),
        (getattr(fy, "fill_type", None), getattr(getattr(fy, "fgColor", None), "rgb", None)),
        tuple((getattr(s, "style", None), _farge(s)) for s in
              (getattr(r, "left", None), getattr(r, "right", None), getattr(r, "top", None), getattr(r, "bottom", None))),
        (getattr(j, "horizontal", None), getattr(j, "vertical", None), getattr(j, "wrap_text", None)),
        c.number_format,
    )


def celler(a: str, b: str, *, bare_verdier: bool, fra_rad: int) -> list[str]:
    import openpyxl

    ut = []
    wa, wb = openpyxl.load_workbook(os_sti(a)), openpyxl.load_workbook(os_sti(b))
    if wa.sheetnames != wb.sheetnames:
        ut.append(f"arknavn: A {wa.sheetnames} B {wb.sheetnames}")
    for navn in wa.sheetnames:
        if navn not in wb.sheetnames:
            continue
        x, y = wa[navn], wb[navn]
        maks_r = max(x.max_row, y.max_row)
        maks_k = max(x.max_column, y.max_column)
        for r in range(fra_rad, maks_r + 1):
            for k in range(1, maks_k + 1):
                ca, cb = x.cell(r, k), y.cell(r, k)
                if ca.value != cb.value:
                    ut.append(f"{navn}!{ca.coordinate}: verdi A {ca.value!r} B {cb.value!r}")
                elif not bare_verdier and _stil(ca) != _stil(cb):
                    ut.append(f"{navn}!{ca.coordinate}: stil A {_stil(ca)} B {_stil(cb)}")
        if bare_verdier:
            continue
        for r in range(1, maks_r + 1):
            if x.row_dimensions[r].height != y.row_dimensions[r].height:
                ut.append(f"{navn} rad {r}: høyde A {x.row_dimensions[r].height} B {y.row_dimensions[r].height}")
        for kol in sorted(set(x.column_dimensions) | set(y.column_dimensions)):
            if x.column_dimensions[kol].width != y.column_dimensions[kol].width:
                ut.append(f"{navn} kolonne {kol}: bredde A {x.column_dimensions[kol].width} B {y.column_dimensions[kol].width}")
        if sorted(map(str, x.merged_cells.ranges)) != sorted(map(str, y.merged_cells.ranges)):
            ut.append(f"{navn}: flettede celler ulike")
        cfa = [(str(c.sqref), [(r.operator, r.formula, str(r.dxf)) for r in c.rules]) for c in x.conditional_formatting]
        cfb = [(str(c.sqref), [(r.operator, r.formula, str(r.dxf)) for r in c.rules]) for c in y.conditional_formatting]
        if cfa != cfb:
            ut.append(f"{navn}: betinget formatering ulik")
        ia = [(i.anchor._from.col, i.anchor._from.row, i.width, i.height) for i in x._images]
        ib = [(i.anchor._from.col, i.anchor._from.row, i.width, i.height) for i in y._images]
        if ia != ib:
            ut.append(f"{navn}: bilder A {ia} B {ib}")
        for felt in ("freeze_panes",):
            if getattr(x, felt) != getattr(y, felt):
                ut.append(f"{navn}: {felt} A {getattr(x, felt)} B {getattr(y, felt)}")
        if x.auto_filter.ref != y.auto_filter.ref:
            ut.append(f"{navn}: filter A {x.auto_filter.ref} B {y.auto_filter.ref}")
        if (x.page_setup.orientation, x.print_title_rows) != (y.page_setup.orientation, y.print_title_rows):
            ut.append(f"{navn}: utskrift ulik")
    return ut


def main(argv=None) -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("a")
    ap.add_argument("b")
    ap.add_argument("--verdier", action="store_true", help="bare verdier og formeltekst")
    ap.add_argument("--fra-rad", type=int, default=1)
    a = ap.parse_args(argv)
    forskjeller = []
    if not a.verdier:
        forskjeller += zip_deler(a.a, a.b)
    forskjeller += celler(a.a, a.b, bare_verdier=a.verdier, fra_rad=a.fra_rad)
    for f in forskjeller[:60]:
        print("  " + f)
    if len(forskjeller) > 60:
        print(f"  … og {len(forskjeller) - 60} til")
    if forskjeller:
        print(f"ULIKT: {len(forskjeller)} forskjell(er)")
        return 1
    print("identisk" + (" (verdier og formler" + (f" fra rad {a.fra_rad}" if a.fra_rad > 1 else "") + ")" if a.verdier
                        else " (alle deler unntatt docProps, celler, stil, mål, CF, bilder)"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
