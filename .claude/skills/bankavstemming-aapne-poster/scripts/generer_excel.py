"""
Genererer Excel-rapport 'Åpne poster i bankavstemming' med Silfer-header.

Bruk:
    python generer_excel.py <data.json> <output.xlsx>

Se SKILL.md for datastrukturen. Skriptet er klient-uavhengig – alt klient-
spesifikt (selskap, kontoer, poster) kommer fra data.json.

Poster kan oppgis strukturert ({"dato","tekst","belop","kilde"}) ELLER som rå
uttrekksstrenger ({"raw": "2026-06-12 −590,00 ZETTLE...", "kilde": "Bankside"}).
Skriptet normaliserer råstrenger selv (håndterer − / mellomrom-tusenskille /
komma-desimal og «| Detaljer:»).
"""
import json, sys, os, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage

DARK, GREEN, LIGHT, BORDER, GREY = "1B2E28", "2D5A4A", "F7F9F8", "D4DDD8", "6B7280"
NUMFMT = '#,##0.00;[Red]-#,##0.00'
LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "silfer_logo.png")
MND = ["Januar", "Februar", "Mars", "April", "Mai", "Juni", "Juli", "August",
       "September", "Oktober", "November", "Desember"]

# Poster som IKKE krever dokumentasjon fra kunden – filtreres bort fra rapporten.
# «Fylt på» = interne kortpåfyll fra driftskontoen; gebyr = bank-/kortgebyr.
IRRELEVANT_MØNSTRE = ["fylt på", "fyll på", "gebyr"]


def er_irrelevant(tekst):
    t = (tekst or "").lower()
    return any(m in t for m in IRRELEVANT_MØNSTRE)


# ── Tekst-rensing ──────────────────────────────────────────────────────────
# Gjør rå bankbeskrivelser korte og lesbare for kunden. Eksempel:
#   "VisaVALUTAKURS: 11,0353 – EUR 126,50 GOOGLE:WORKSPACE E, 05469769..."
#   → "Google Workspace 126,50 EUR"
# Kjøres kun på råstrenger ("raw"); eksplisitt "tekst" i data.json røres ikke.
BETALINGSPREFIKS = (r"^\s*(?:Visa|Giro|Avtalegiro|Straksbetaling|Straksoverf\w+|"
                    r"Nettgiro|Nettbank|MobilePay|Vipps|Kort|Bankaxept)\s*")
VALUTA = "EUR|USD|GBP|SEK|DKK|CHF|PLN|JPY|CAD|AUD|NOK"
# Ord som skal beholde store bokstaver (selskapsformer, valuta, forkortelser)
# Selskapsformer/forkortelser som skal beholde store bokstaver. Anglo-former
# (Inc/Ltd/LLC/GmbH) utelates med vilje – de leses penere som «Inc.», «Ltd».
ACRONYMER = {"AS", "ASA", "A/S", "ANS", "DA", "BA", "SA", "NUF", "AB", "OY",
             "EUR", "USD", "GBP", "SEK", "DKK", "NOK", "CHF", "PLN", "JPY",
             "UK", "US", "EU", "IT", "AI", "VVS"}


def _cap_ord(w):
    up = w.upper()
    if up in ACRONYMER:
        return up
    if any(ch.isdigit() for ch in w):     # koder som 26_247480 – la stå
        return w
    if w.isupper() or w.islower():         # KUN helt caps/lower gjøres pene
        return w.capitalize()
    return w                               # blandet case (iZettle) – la stå


def clean_tekst(tekst):
    if not tekst:
        return tekst
    t = tekst.strip()
    # 1) fjern betalingstype-prefiks limt foran (VisaVALUTAKURS, GiroSTOREBRAND)
    t = re.sub(BETALINGSPREFIKS, "", t, flags=re.I)
    # 2) fjern valutakurs-støy
    t = re.sub(r"VALUTAKURS:?\s*[\d.,]+", " ", t, flags=re.I)
    # 3) trekk ut utenlandsk valutabeløp (EUR 126,50) og legg som suffiks til slutt
    cur = ""
    m = re.search(r"\b(" + VALUTA + r")\s+(\d[\d .]*,\d{2})", t, flags=re.I)
    if m:
        cur = f"{m.group(2).strip()} {m.group(1).upper()}"
        t = t.replace(m.group(0), " ")
    # 4) kutt bort merkelapp-haler (bilag/betaling/faktura/ref/detaljer m.m.)
    t = re.sub(r"(Bilag:|Betaling:|Betalingstype:|Detaljer:|Beskrivelse:|Annet:|"
               r"Melding:|KID:|Fakturanr|Faktura\s*nummer|Ref[:.]|Til\s+konto|"
               r"Fra\s+konto).*$", " ", t, flags=re.I)
    # 5) fjern norske kontonummer (1234 56 78901 / 1234.56.78901) og lange arkivnr
    t = re.sub(r"\b\d{4}[ .]\d{2}[ .]\d{5}\b", " ", t)
    t = re.sub(r",\s*\d{4,}.*$", " ", t)
    t = re.sub(r"\b\d{7,}\b", " ", t)
    # 6) fjern separatorer og skilletegn-støy
    t = re.sub(r"[|–]", " ", t)
    t = re.sub(r"\s[-]\s", " ", t)
    t = t.replace(":", " ").replace("*", " ").replace("_", " ")
    # 7) rydd whitespace og kant-tegn
    t = re.sub(r"\s+", " ", t).strip(" ,.-")
    # 8) dropp etterhengende ensom stor bokstav (f.eks. «Workspace E»)
    ord_ = t.split()
    if len(ord_) >= 2 and re.fullmatch(r"[A-ZÆØÅ]", ord_[-1]):
        ord_ = ord_[:-1]
    # dropp etterhengende løs preposisjon/konjunksjon (f.eks. «Plasseringskonto Fra»)
    STOPP = {"fra", "til", "og", "for", "med", "av", "på"}
    while len(ord_) >= 2 and ord_[-1].lower() in STOPP:
        ord_ = ord_[:-1]
    t = " ".join(_cap_ord(w) for w in ord_).strip()
    if cur:
        t = (t + " " + cur).strip()
    return t or tekst.strip()


def normalize_kilde(k):
    if not k:
        return ""
    s = str(k).strip().lower()
    if s.startswith("bank") or s == "right" or "kontoutskrift" in s:
        return "Bankside"
    if s.startswith("tripletex") or s == "left" or "bokf" in s:
        return "Tripletex-side"
    return str(k)


def parse_amount(txt):
    """'−5 060,50' / '5\u00a0060,50' / '-1.808,25' -> float."""
    t = txt.strip().replace("\u2212", "-").replace("\u00a0", " ")
    neg = t.startswith("-")
    t = t.lstrip("-").replace(" ", "").replace(".", "").replace(",", ".")
    val = float(re.sub(r"[^0-9.]", "", t) or 0)
    return -val if neg else val


def parse_raw(raw):
    """'2026-06-08 5 060,50 Oase ... | Detaljer: Fylt på' -> (dato, belop, tekst)."""
    s = raw.strip().replace("\u00a0", " ")
    # Bel\u00f8pet slutter ved \u00ab,dd\u00bb; stopp der s\u00e5 etterf\u00f8lgende referansenr (lange
    # sifferstrenger) ikke slukes inn i bel\u00f8pet.
    m = re.match(r"^(\d{4}-\d{2}-\d{2})\s+([\u2212\-]?\d[\d .]*,\d{2})\s+(.*)$", s)
    if not m:
        m = re.match(r"^(\d{4}-\d{2}-\d{2})\s+([\u2212\-]?[\d .,]+?)\s+(.*)$", s)
    if not m:
        return None
    dato, amt, tekst = m.group(1), m.group(2), m.group(3).strip()
    tekst = tekst.replace(" | Detaljer:", " –").replace("| Detaljer:", "–").strip()
    return dato, parse_amount(amt), tekst


def norm_post(p):
    """Return dict with dato/tekst/belop/kilde, accepting structured or raw."""
    kilde = normalize_kilde(p.get("kilde", ""))
    if "raw" in p and "belop" not in p:
        parsed = parse_raw(p["raw"])
        if not parsed:
            raise ValueError(f"Kunne ikke tolke rad: {p['raw']!r}")
        dato, belop, tekst = parsed
        # rens råstrengen til lesbar tekst; behold original for filtrering
        return {"dato": dato, "tekst": clean_tekst(tekst), "_match": tekst,
                "belop": round(belop, 2), "kilde": kilde}
    return {"dato": p["dato"], "tekst": p["tekst"], "_match": p["tekst"],
            "belop": round(float(p["belop"]), 2), "kilde": kilde}


def month_sort_key(name):
    n = name.strip().split("–")[0].split("-")[0].strip().capitalize()
    return MND.index(n) if n in MND else 99


def build(data, out_path):
    selskap = data.get("selskap", "Ukjent klient")
    orgnr = data.get("orgnr", "").strip()
    aar = str(data.get("aar", ""))
    dato = data.get("dato", "").strip() or datetime.date.today().strftime("%d.%m.%Y")

    # Standard: ta KUN med bankside-poster (i bank, mangler bilag). Tripletex-side
    # (bokført, ikke sett i bank – ofte timing) tas ikke med med mindre man ber om
    # det via "inkluder_tripletex_side": true i data.json.
    kun_bankside = not data.get("inkluder_tripletex_side", False)

    # Normaliser alle poster, og filtrer bort ikke-relevante (Fylt på / gebyr)
    kontoer = []
    utelatt = 0
    utelatt_tx = 0
    for k in data["kontoer"]:
        maaneder = []
        for m in k.get("maaneder", []):
            poster = []
            for p in m.get("poster", []):
                np = norm_post(p)
                if kun_bankside and np["kilde"] == "Tripletex-side":
                    utelatt_tx += 1
                    continue
                if er_irrelevant(np.get("_match", np["tekst"])):
                    utelatt += 1
                    continue
                poster.append(np)
            maaneder.append({"maaned": m["maaned"], "status": m.get("status", ""), "poster": poster})
        maaneder.sort(key=lambda m: month_sort_key(m["maaned"]))
        kontoer.append({"nummer": str(k["nummer"]), "navn": k.get("navn", ""),
                        "merknad": k.get("merknad", ""), "maaneder": maaneder})

    wb = Workbook()
    thin = Side(style="thin", color=BORDER)
    bottom = Border(bottom=thin)

    def header_block(ws, title, subtitle, ncols):
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 20
        ws.row_dimensions[2].height = 26
        ws.row_dimensions[3].height = 16
        ws.row_dimensions[4].height = 6
        if os.path.exists(LOGO):
            img = XLImage(LOGO); img.width, img.height = 176, 54
            ws.add_image(img, "A1")
        ws["D2"] = title
        ws["D2"].font = Font(name="Arial", size=15, bold=True, color=DARK)
        ws["D2"].alignment = Alignment(vertical="center")
        ws["D3"] = subtitle
        ws["D3"].font = Font(name="Arial", size=9, color=GREY)
        for c in range(1, ncols + 1):
            ws.cell(row=5, column=c).fill = PatternFill("solid", fgColor=DARK)
        ws.row_dimensions[5].height = 3

    def head_cells(ws, hrow, headers, right_idx):
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=hrow, column=i, value=h)
            c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=DARK)
            c.alignment = Alignment(horizontal="right" if i in right_idx else "left", vertical="center")
        ws.row_dimensions[hrow].height = 20

    sub = f"{selskap}   |   " + (f"Org.nr. {orgnr}   |   " if orgnr else "") + f"{dato}"

    # ───── Ark: Åpne poster ─────
    ws = wb.active
    ws.title = "Åpne poster"
    header_block(ws, f"Åpne poster i bankavstemming {aar}", sub, 9)
    HROW = 7
    head_cells(ws, HROW, ["Konto", "Kontonavn", "Måned", "Dato", "Tekst",
                          "Beløp (kr)", "Kilde", "Dokumentasjon", "Status"], {6})
    row = HROW + 1
    for k in kontoer:
        for m in k["maaneder"]:
            for p in m["poster"]:
                try:
                    dt = datetime.date.fromisoformat(p["dato"])
                except Exception:
                    dt = p["dato"]
                vals = [int(k["nummer"]) if k["nummer"].isdigit() else k["nummer"],
                        k["navn"], m["maaned"], dt, p["tekst"], p["belop"], p["kilde"], "", ""]
                for i, v in enumerate(vals, 1):
                    c = ws.cell(row=row, column=i, value=v)
                    c.font = Font(name="Arial", size=9, color="1A1A1A")
                    c.border = bottom
                    if i == 6:
                        c.number_format = NUMFMT; c.alignment = Alignment(horizontal="right")
                    elif i == 4 and isinstance(dt, datetime.date):
                        c.number_format = "yyyy-mm-dd"
                    if row % 2 == 0:
                        c.fill = PatternFill("solid", fgColor=LIGHT)
                row += 1
    last = row - 1
    if last >= HROW + 1:
        srow = row + 1
        ws.cell(row=srow, column=5, value="Sum åpne poster").font = Font(name="Arial", size=9, bold=True, color=DARK)
        sc = ws.cell(row=srow, column=6, value=f"=SUM(F{HROW+1}:F{last})")
        sc.font = Font(name="Arial", size=9, bold=True, color=DARK)
        sc.number_format = NUMFMT; sc.alignment = Alignment(horizontal="right")
        cc = ws.cell(row=srow, column=1, value=f"=COUNTA(E{HROW+1}:E{last})")
        cc.font = Font(name="Arial", size=9, bold=True, color=DARK)
        ws.auto_filter.ref = f"A{HROW}:I{last}"
    for col, w in {"A": 8, "B": 24, "C": 10, "D": 12, "E": 34, "F": 13, "G": 14, "H": 26, "I": 14}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{HROW+1}"

    # ───── Ark: Sammendrag ─────
    ws2 = wb.create_sheet("Sammendrag")
    header_block(ws2, "Sammendrag – åpne poster per konto", f"{selskap}   |   {aar}", 5)
    H2 = 7
    head_cells(ws2, H2, ["Konto", "Kontonavn", "Åpne poster", "Sum beløp (kr)", "Merknad"], {3, 4})
    r = H2 + 1
    data_first, data_last = HROW + 1, max(last, HROW + 1)
    for k in kontoer:
        knr = int(k["nummer"]) if k["nummer"].isdigit() else k["nummer"]
        n_poster = sum(len(m["poster"]) for m in k["maaneder"])
        merk = k["merknad"] or ("Ingen åpne poster" if n_poster == 0 else "")
        vals = [knr, k["navn"],
                f"=COUNTIF('Åpne poster'!$A${data_first}:$A${data_last},A{r})",
                f"=SUMIF('Åpne poster'!$A${data_first}:$A${data_last},A{r},'Åpne poster'!$F${data_first}:$F${data_last})",
                merk]
        for i, v in enumerate(vals, 1):
            c = ws2.cell(row=r, column=i, value=v)
            c.font = Font(name="Arial", size=9, color="1A1A1A")
            c.border = bottom
            if i == 4:
                c.number_format = NUMFMT; c.alignment = Alignment(horizontal="right")
            if i == 3:
                c.alignment = Alignment(horizontal="right")
            if r % 2 == 0:
                c.fill = PatternFill("solid", fgColor=LIGHT)
        r += 1
    ws2.cell(row=r, column=2, value="Totalt").font = Font(name="Arial", size=9, bold=True, color=DARK)
    tc = ws2.cell(row=r, column=3, value=f"=SUM(C{H2+1}:C{r-1})")
    tc.font = Font(name="Arial", size=9, bold=True, color=DARK); tc.alignment = Alignment(horizontal="right")
    ts = ws2.cell(row=r, column=4, value=f"=SUM(D{H2+1}:D{r-1})")
    ts.font = Font(name="Arial", size=9, bold=True, color=DARK)
    ts.number_format = NUMFMT; ts.alignment = Alignment(horizontal="right")
    for col, w in {"A": 8, "B": 24, "C": 13, "D": 15, "E": 52}.items():
        ws2.column_dimensions[col].width = w
    tx_note = ("Kun bankside-poster er tatt med (i bank, mangler bilag). "
               if kun_bankside else
               "Både bank- og Tripletex-side er tatt med. ")
    note = ("Kilde: Automatisk bankavstemming i Tripletex. Åpne (uavstemte) poster "
            + (data.get("periode", "") or aar)
            + ". " + tx_note
            + "Konto for skattetrekk (typ. 1950) er holdt utenfor. «Fylt på» "
            "(kortpåfyll) og bank-/kortgebyr er utelatt – disse krever ingen "
            "dokumentasjon.")
    ws2.cell(row=r + 2, column=1, value=note).font = Font(name="Arial", size=8, italic=True, color=GREY)

    wb.move_sheet("Sammendrag", -1)  # Sammendrag først
    wb.save(out_path)
    antall = last - HROW if last >= HROW + 1 else 0
    print(f"Lagret {out_path} · {antall} poster · {len(kontoer)} kontoer · "
          f"{utelatt} utelatt (Fylt på / gebyr) · {utelatt_tx} Tripletex-side utelatt")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Bruk: python generer_excel.py <data.json> <output.xlsx>"); sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        build(json.load(f), sys.argv[2])
