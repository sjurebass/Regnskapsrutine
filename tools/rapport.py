"""Excel-rapporter med Silfer-profil.

Stilen er hentet fra lønnsavstemmingene som allerede ligger i kundemappene:
Arial, mørkegrønn header (#1B2E28), seksjonsrad (#2D5A4A), grå undertittel.
"""

from __future__ import annotations

import os
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .modeller import GUL, INFO, ROD, Funn, Periode, Resultat

MORK = "FF1B2E28"
GRONN = "FF2D5A4A"
GRA = "FF6B7280"
SVART = "FF1A1A1A"
LYS = "FFF3F4F6"
FARGE_STATUS = {ROD: "FFC0392B", GUL: "FFB7791F", INFO: GRA, "grønn": GRONN, "feil": "FF7C3AED"}

TYNN = Side(style="thin", color="FFD1D5DB")
RAMME = Border(bottom=TYNN)

LOGO = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "mal", "silfer-logo.png")


def _tittelblokk(ws, tittel: str, undertittel: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 19.5
    ws.row_dimensions[2].height = 25.5
    ws.row_dimensions[3].height = 15.75
    ws.row_dimensions[4].height = 6.0
    ws.row_dimensions[5].height = 3.0
    ws["D2"] = tittel
    ws["D2"].font = Font(name="Arial", size=15, bold=True, color=MORK)
    ws["D3"] = undertittel
    ws["D3"].font = Font(name="Arial", size=9, color=GRA)
    if os.path.exists(LOGO):
        try:
            bilde = XLImage(LOGO)
            forhold = (bilde.width or 120) / (bilde.height or 46)
            bilde.height = 46
            bilde.width = int(46 * forhold)
            ws.add_image(bilde, "A1")
        except Exception:
            pass


def _tabellhode(ws, rad: int, kolonner: list[tuple[str, int]]) -> None:
    ws.row_dimensions[rad].height = 19.5
    for i, (navn, bredde) in enumerate(kolonner, start=1):
        celle = ws.cell(row=rad, column=i, value=navn)
        celle.font = Font(name="Arial", size=9, bold=True, color="FFFFFFFF")
        celle.fill = PatternFill("solid", fgColor=MORK)
        celle.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = bredde


def _seksjonsrad(ws, rad: int, tekst: str, antall_kolonner: int) -> None:
    ws.row_dimensions[rad].height = 18
    for i in range(1, antall_kolonner + 1):
        celle = ws.cell(row=rad, column=i, value=tekst if i == 1 else None)
        celle.font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
        celle.fill = PatternFill("solid", fgColor=GRONN)
        celle.alignment = Alignment(vertical="center")


def _skriv(ws, rad: int, verdier: list, fet: bool = False, farge: str | None = None) -> None:
    for i, verdi in enumerate(verdier, start=1):
        celle = ws.cell(row=rad, column=i, value=verdi)
        celle.font = Font(name="Arial", size=9, bold=fet, color=farge or SVART)
        celle.alignment = Alignment(vertical="top", wrap_text=(i >= 6))
        celle.border = RAMME
        if isinstance(verdi, (int, float)):
            celle.number_format = "# ##0;-# ##0"
            celle.alignment = Alignment(vertical="top", horizontal="right")


def _funnark(wb, resultat: Resultat) -> None:
    ws = wb.create_sheet("Funn")
    _tittelblokk(
        ws,
        f"Månedskontroll {resultat.periode}",
        f"{resultat.klient.navn}   |   {resultat.periode.navn}   |   "
        f"kjørt {date.today().strftime('%d.%m.%Y')}",
    )
    kolonner = [
        ("Kontroll", 10),
        ("Område", 12),
        ("Status", 9),
        ("Konto", 12),
        ("Beløp", 14),
        ("Funn", 52),
        ("Detaljer", 70),
        ("Nøkkel", 12),
    ]
    _tabellhode(ws, 7, kolonner)
    ws.freeze_panes = "A8"
    rad = 8
    omraade = None
    for f in resultat.funn:
        if f.omraade != omraade:
            omraade = f.omraade
            _seksjonsrad(ws, rad, omraade, len(kolonner))
            rad += 1
        status = "kjent" if f.kjent else f.alvorlighet
        detalj = f.detalj
        if f.kjent and f.kjent_kommentar:
            detalj = f"KVITTERT: {f.kjent_kommentar}. {detalj}"
        _skriv(
            ws,
            rad,
            [f.kontroll, f.omraade, status, f.konto or "", f.belop, f.tittel, detalj, f.noekkel],
            farge=GRA if f.kjent else FARGE_STATUS.get(f.alvorlighet),
        )
        rad += 1
    if not resultat.funn:
        _skriv(ws, rad, ["", "", "grønn", "", None, "Ingen avvik funnet", "", ""], fet=True)


def _sammendragsark(wb, resultat: Resultat, advarsler: list[str]) -> None:
    ws = wb.create_sheet("Sammendrag", 0)
    _tittelblokk(
        ws,
        f"Månedskontroll {resultat.periode}",
        f"{resultat.klient.navn}   |   {resultat.periode.navn}   |   "
        f"kjørt {datetime.now().strftime('%d.%m.%Y %H:%M')}",
    )
    kolonner = [("Post", 34), ("Verdi", 52), ("", 16), ("", 14), ("", 14), ("", 30), ("", 40)]
    _tabellhode(ws, 7, kolonner)
    rad = 8
    apne = resultat.aapne_funn
    teller = {
        "Røde avvik": sum(1 for f in apne if f.alvorlighet == ROD),
        "Gule avvik": sum(1 for f in apne if f.alvorlighet == GUL),
        "Til informasjon": sum(1 for f in apne if f.alvorlighet == INFO),
        "Kvittert som kjent": sum(1 for f in resultat.funn if f.kjent),
    }
    _skriv(ws, rad, ["Status", resultat.status.upper()], fet=True, farge=FARGE_STATUS.get(resultat.status))
    rad += 1
    for navn, antall in teller.items():
        _skriv(ws, rad, [navn, antall])
        rad += 1
    _skriv(ws, rad, ["Kontroller kjørt", ", ".join(resultat.kjorte) or "-"])
    rad += 2

    viktigste = [f for f in apne if f.alvorlighet in (ROD, GUL)][:10]
    if viktigste:
        _seksjonsrad(ws, rad, "Det som trenger oppmerksomhet", len(kolonner))
        rad += 1
        for i, f in enumerate(viktigste, start=1):
            _skriv(
                ws,
                rad,
                [f"{i}. {f.kontroll}", f.tittel, f.belop, "", "", f.konto or "", f.detalj[:200]],
                farge=FARGE_STATUS.get(f.alvorlighet),
            )
            rad += 1
        rad += 1

    if resultat.hoppet_over or resultat.feil or advarsler:
        _seksjonsrad(ws, rad, "Kjøringsnotater", len(kolonner))
        rad += 1
        for kid, grunn in resultat.feil:
            _skriv(ws, rad, [f"Feil i {kid}", grunn], farge=FARGE_STATUS["feil"])
            rad += 1
        for kid, grunn in resultat.hoppet_over:
            _skriv(ws, rad, [f"Hoppet over {kid}", grunn], farge=GRA)
            rad += 1
        for a in advarsler[:10]:
            _skriv(ws, rad, ["Advarsel", a[:180]], farge=GRA)
            rad += 1


def _saldoark(wb, resultat: Resultat, saldi_perioden: list[dict], saldi_hittil: list[dict]) -> None:
    ws = wb.create_sheet("Saldobalanse")
    _tittelblokk(
        ws,
        "Saldobalanse",
        f"{resultat.klient.navn}   |   {resultat.periode.navn} og hittil i år",
    )
    kolonner = [
        ("Konto", 10),
        ("Navn", 40),
        ("IB året", 16),
        ("Bevegelse mnd", 16),
        ("Bevegelse hittil", 18),
        ("UB", 16),
        ("", 10),
    ]
    _tabellhode(ws, 7, kolonner)
    ws.freeze_panes = "A8"
    mnd = {}
    for r in saldi_perioden:
        nummer = (r.get("account") or {}).get("number")
        if nummer is not None:
            mnd[str(nummer)] = float(r.get("sumAmount") or 0)
    rad = 8
    for r in sorted(saldi_hittil, key=lambda x: str((x.get("account") or {}).get("number") or "")):
        konto = r.get("account") or {}
        nummer = str(konto.get("number") or "")
        ib = float(r.get("openingBalance") or 0)
        hittil = float(r.get("sumAmount") or 0)
        ub_ = float(r.get("closingBalance") or 0)
        if abs(ib) < 0.5 and abs(hittil) < 0.5 and abs(ub_) < 0.5:
            continue
        _skriv(ws, rad, [nummer, konto.get("name") or "", ib, mnd.get(nummer, 0.0), hittil, ub_])
        rad += 1


def _lagre(wb, sti: str) -> str:
    """Lagrer, med retry hvis filen er låst (åpen i Excel eller skannet av
    antivirus). Gir opp til et alternativt filnavn framfor å miste kjøringen."""
    import time

    for forsok in range(3):
        try:
            wb.save(sti)
            return sti
        except PermissionError:
            if forsok < 2:
                time.sleep(1.5)
                continue
            rot, ext = os.path.splitext(sti)
            alternativ = f"{rot} ({datetime.now().strftime('%H%M%S')}){ext}"
            wb.save(alternativ)
            print(f"  ! {os.path.basename(sti)} var låst – lagret som {os.path.basename(alternativ)}")
            return alternativ
    return sti


def skriv_klientrapport(
    resultat: Resultat,
    saldi_perioden: list[dict],
    saldi_hittil: list[dict],
    advarsler: list[str],
    mappe: str,
) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    _funnark(wb, resultat)
    _sammendragsark(wb, resultat, advarsler)
    _saldoark(wb, resultat, saldi_perioden, saldi_hittil)
    os.makedirs(mappe, exist_ok=True)
    sti = os.path.join(mappe, f"Kontroll {resultat.periode} {resultat.klient.navn}.xlsx")
    return _lagre(wb, sti)


def skriv_statusrapport(resultater: list[Resultat], periode: Periode, mappe: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Status"
    _tittelblokk(
        ws,
        f"Statusrapport {periode}",
        f"Alle klienter i kontrollrutinen   |   {periode.navn}   |   "
        f"kjørt {datetime.now().strftime('%d.%m.%Y %H:%M')}",
    )
    kolonner = [
        ("Klient", 32),
        ("Status", 10),
        ("Rød", 7),
        ("Gul", 7),
        ("Info", 7),
        ("Viktigste funn", 80),
        ("Rapport", 46),
    ]
    _tabellhode(ws, 7, kolonner)
    ws.freeze_panes = "A8"
    rad = 8
    rekkefolge = {"feil": 0, ROD: 1, GUL: 2, "grønn": 3}
    for r in sorted(resultater, key=lambda x: (rekkefolge.get(x.status, 9), x.klient.navn)):
        apne = r.aapne_funn
        topp = "; ".join(
            f"{f.kontroll} {f.tittel}" for f in apne if f.alvorlighet in (ROD, GUL)
        )[:400]
        if r.feil and not topp:
            topp = "Kjøringsfeil: " + "; ".join(f"{k}: {m}" for k, m in r.feil)[:300]
        _skriv(
            ws,
            rad,
            [
                r.klient.navn,
                r.status.upper(),
                sum(1 for f in apne if f.alvorlighet == ROD),
                sum(1 for f in apne if f.alvorlighet == GUL),
                sum(1 for f in apne if f.alvorlighet == INFO),
                topp or "Ingen avvik",
                f"{r.klient.mappe}/{periode.aar}/Kontroll {periode} {r.klient.navn}.xlsx",
            ],
            farge=FARGE_STATUS.get(r.status),
        )
        rad += 1
    os.makedirs(mappe, exist_ok=True)
    sti = os.path.join(mappe, f"{periode} Statusrapport alle klienter.xlsx")
    return _lagre(wb, sti)
